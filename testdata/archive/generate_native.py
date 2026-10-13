"""Optional independent-producer controls; use bundled Python with docx/openpyxl.
No parser/runtime dependency. ZIP timestamps and core dates are normalized.
"""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
from datetime import datetime
from importlib.metadata import version
import re
from docx import Document
from openpyxl import Workbook
ROOT=Path(__file__).parent

def save(name,raw):
    with ZipFile(BytesIO(raw)) as source, ZipFile(ROOT/name,'w') as target:
        for key in sorted(source.namelist()):
            info=ZipInfo(key,(2020,1,1,0,0,0)); info.compress_type=ZIP_DEFLATED; info.external_attr=0o100644<<16
            data=source.read(key)
            if key=='docProps/core.xml':
                data=re.sub(rb'(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)',rb'\g<1>2020-01-01T00:00:00Z\g<2>',data)
            target.writestr(info,data)

doc=Document();doc.core_properties.created=datetime(2020,1,1);doc.core_properties.modified=datetime(2020,1,1)
doc.add_paragraph('虚构健康报告：患者张明 出生1980-03-12')
table=doc.add_table(rows=2,cols=2)
for row,values in zip(table.rows,[('项目','结果'),('医生','周青')]):
    for cell,value in zip(row.cells,values):cell.text=value
doc.add_paragraph('备注：此文件仅用于合成测试')
b=BytesIO();doc.save(b);save('native-health.docx',b.getvalue())
book=Workbook();book.properties.created=datetime(2020,1,1);book.properties.modified=datetime(2020,1,1)
sheet=book.active;sheet.title='家庭设备';sheet.append(['地址','测试苑1号测试室']);sheet.append(['型号','DEMO-AC']);sheet.append(['保修到期',datetime(2028,5,31)]);sheet['B3'].number_format='yyyy-mm-dd'
sheet.merge_cells('A5:B5');sheet['A5']='示例家庭设备清单';sheet['C7']=0;sheet.row_dimensions[7].hidden=True
b=BytesIO();book.save(b);save('native-home.xlsx',b.getvalue())
(ROOT/'native-producers.txt').write_text('python-docx='+version('python-docx')+'\nopenpyxl='+version('openpyxl')+'\n')
