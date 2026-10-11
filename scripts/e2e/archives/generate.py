#!/usr/bin/env python3
"""Create only fictional source files. Expected answers are never sent to workers."""
import argparse, hashlib, json, shutil, re
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED, ZipInfo
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader
from pypdf import PdfReader, PdfWriter
from docx import Document
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[3]
pdfmetrics.registerFont(TTFont('ArchiveSyntheticEmbedded','/System/Library/Fonts/Supplemental/Arial Unicode.ttf'))
FONT = '/System/Library/Fonts/Supplemental/Songti.ttc'

def image(lines):
    im=Image.new('RGB',(1600,1100),'white'); d=ImageDraw.Draw(im); f=ImageFont.truetype(FONT,44)
    for i,line in enumerate(lines): d.text((90,80+i*90),line,font=f,fill='black')
    return im

def pdf(path,pages,scan=False):
    c=canvas.Canvas(str(path),pagesize=(800,550),invariant=1)
    for lines in pages:
        if scan: c.drawImage(ImageReader(image(lines)),0,0,width=800,height=550)
        else:
            c.setFont('ArchiveSyntheticEmbedded',20)
            for i,line in enumerate(lines): c.drawString(45,500-i*40,line)
        c.showPage()
    c.save()

def generate(out):
    out.mkdir(parents=True,exist_ok=True)
    files=[]
    def add(id,name,private=False,kind='file'):
        p=out/name
        files.append(dict(id=id,file=name,kind=kind,private=private,sha256=hashlib.sha256(p.read_bytes()).hexdigest(),bytes=p.stat().st_size))
    h1=[['虚构体检报告','患者：张明','出生日期：1980-03-12','报告日期：2026-06-15','医生：周青'],['虚构体检报告','患者：张明  出生日期：1980-03-12','报告日期：2026-06-15','血红蛋白：142 g/L','参考范围：130–175 g/L','医生：周青']]
    pdf(out/'H01.pdf',h1);add('H01','H01.pdf')
    pdf(out/'H02.pdf',[['虚构体检报告','患者：张明','出生日期：1992-08-20','报告日期：2026-07-02','血红蛋白：151 g/L']],True);add('H02','H02.pdf')
    image(['虚构体检报告','患者：张明','出生日期：1980-03-12','报告日期：2026-09-10','血红蛋白：145 g/L']).save(out/'H03.png');add('H03','H03.png',True)
    for id in ['H04a','H04b']:
        (out/(id+'.txt')).write_text('喂奶记录\n宝宝：小禾\n出生日期：2026-06-01\n记录日期：2026-10-03\n奶量：70 ml\n',encoding='utf-8');add(id,id+'.txt',kind='text')
    d=Document();d.add_heading('虚构学生成绩报告',0);d.add_paragraph('学生：李妍\n出生日期：2012-11-06\n学期：2026春季\n学校：示例学校')
    table=d.add_table(rows=1,cols=3)
    for cell,text in zip(table.rows[0].cells,['科目','分数','满分']):cell.text=text
    for row in [('数学','88','100'),('语文','91','100')]:
        for cell,text in zip(table.add_row().cells,row):cell.text=text
    d.save(out/'E01.docx');add('E01','E01.docx')
    w=Workbook();s=w.active;s.title='成绩';s.merge_cells('A1:C1');s['A1']='学生成绩报告';s.append(['学生：李妍 出生日期：2012-11-06']);s.append(['学期：2026秋季']);s.append(['科目','分数','满分']);s.append(['数学',132,150]);w.save(out/'E02.xlsx');add('E02','E02.xlsx')
    (out/'E03.txt').write_text('家庭教育资料\n示例学校运动会\n时间：2026年10月\n',encoding='utf-8');add('E03','E03.txt',kind='text')
    # Keep the old ambiguous parser fixture as a separately hashed negative
    # control. The positive source explicitly labels roles/currency/dates; no
    # classification or extracted answer is passed to the worker.
    # A normal business document, separate from the parser stress workbook.
    # Store the publisher's formula cache explicitly; the parser must not execute it.
    w=Workbook();sheet=w.active;sheet.title='费用'
    sheet.append(['账单编号','费用','收款方','币种','日期'])
    sheet.append(['TEST-BILL-001',128.10,'示例物业','CNY','2026-10-01'])
    sheet.append(['TEST-BILL-002',71.90,'示例物业','CNY','2026-10-01'])
    sheet.append(['合计','=SUM(B2:B3)',None,'CNY'])
    for row in [2,3,4]:sheet[f'B{row}'].number_format='0.00'
    data=BytesIO();w.save(data)
    for id,cached in [('F01',True),('F01-missing',False)]:
        with ZipFile(BytesIO(data.getvalue())) as source, ZipFile(out/(id+'.xlsx'),'w',ZIP_DEFLATED) as target:
            for member in source.infolist():
                raw=source.read(member.filename)
                if member.filename=='xl/worksheets/sheet1.xml':
                    text=raw.decode();old='<f>SUM(B2:B3)</f><v></v>'
                    assert old in text
                    text=text.replace(old,'<f>SUM(B2:B3)</f>'+('<v>200</v>' if cached else ''))
                    for cell,value in [('B2','128.10'),('B3','71.90')]:
                        text,n=re.subn(r'(<c\b[^>]*\br="'+cell+r'"[^>]*><v>)[^<]+(</v>)',lambda m:m[1]+value+m[2],text);assert n==1
                    raw=text.encode()
                target.writestr(member,raw)
        add(id,id+'.xlsx')
    shutil.copyfile(ROOT/'testdata/archive/f01.xlsx',out/'F01-unlabelled-negative.xlsx');add('F01-unlabelled-negative','F01-unlabelled-negative.xlsx')
    (out/'F02.csv').write_text('凭证类型,付款所引用账单编号,付款,币种,日期,收款方\n付款凭证,TEST-BILL-001,128.10,CNY,2026-10-01,示例物业\n',encoding='utf-8');add('F02','F02.csv')
    (out/'F03.csv').write_text('账单编号,收款方,日期,费用,币种\nTEST-BILL-003,示例物业,2026-10-02,20,USD\n',encoding='utf-8');add('F03','F03.csv')
    (out/'F04.txt').write_text('独立物业账单\n账单编号：TEST-BILL-004\n收款方：示例物业\n日期：2026-10-03\n费用：50 CNY\n',encoding='utf-8');add('F04','F04.txt',True,'text')
    im=image(['家庭设备保修资料','房屋地址：测试苑1号测试室','设备：空调','设备型号：DEMO-AC','保修到期：2028-05-31'])
    im.save(out/'P01.jpeg',quality=95);add('P01','P01.jpeg'); im.rotate(90,expand=True).save(out/'P01-rotated.jpeg',quality=95);add('P01-rotated','P01-rotated.jpeg')
    image(['家庭设备维修单','房屋地址：测试苑1号测试室','设备型号：DEMO-AC','维修日期：2026-08-11','维修内容：更换滤网','工程师：陈工']).save(out/'P02.webp',lossless=True);add('P02','P02.webp')
    pdf(out/'P03.pdf',[['虚构房屋使用合同','房屋地址：测试苑1号测试室','业主：张明  出生日期：1980-03-12','居住人：李妍  出生日期：2012-11-06','合同日期：2026-08-01']]);add('P03','P03.pdf')
    (out/'A01.txt').write_text('张明的体检报告到了\n',encoding='utf-8');add('A01','A01.txt',kind='text')
    wr=PdfWriter();wr.append(PdfReader(out/'H01.pdf'));wr.encrypt('synthetic-password');wr.write(out/'A02-encrypted.pdf');add('A02-encrypted','A02-encrypted.pdf')
    (out/'A02-corrupt.docx').write_bytes(b'PK\x03\x04synthetic truncated zip');add('A02-corrupt','A02-corrupt.docx')
    (out/'A02-unknown.bin').write_bytes(b'ATRIUM SYNTHETIC UNKNOWN FORMAT\x00\xff');add('A02-unknown','A02-unknown.bin')
    (out/'manifest.json').write_text(json.dumps({'synthetic_only':True,'fixture_version':'t15-v5','fixtures':files},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'generated':len(files),'directory':str(out)}))

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('directory',type=Path);a=p.parse_args();generate(a.directory)
